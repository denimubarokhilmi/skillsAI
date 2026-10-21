#!/usr/bin/env python3
"""Polish a basic XLSX workbook so it is presentable by default."""

from __future__ import annotations

import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

HEADER_FILL = "1F4E78"
HEADER_FONT = "FFFFFF"
BORDER = Border(bottom=Side(style="thin", color="D9E2F3"))
BODY_BORDER = Border(left=Side(style="thin", color="E7E6E6"), right=Side(style="thin", color="E7E6E6"), top=Side(style="thin", color="E7E6E6"), bottom=Side(style="thin", color="E7E6E6"))


def used_range(ws):
    return ws.max_row, ws.max_column


def header_row(ws):
    for row in range(1, min(ws.max_row, 10) + 1):
        values = [ws.cell(row, col).value for col in range(1, ws.max_column + 1)]
        filled = sum(value not in (None, "") for value in values)
        if filled >= max(2, ws.max_column // 3):
            return row
    return 1


def looks_numeric(value):
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def set_number_formats(ws, header):
    for col in range(1, ws.max_column + 1):
        title = str(ws.cell(header, col).value or "").lower()
        for row in range(header + 1, ws.max_row + 1):
            cell = ws.cell(row, col)
            if cell.value is None:
                continue
            if any(key in title for key in ["%", "rate", "margin", "growth", "ratio"]):
                cell.number_format = "0.0%"
            elif any(key in title for key in ["amount", "revenue", "cost", "price", "sales", "profit", "budget"]):
                cell.number_format = '#,##0;(#,##0);-'
            elif any(key in title for key in ["date", "tanggal"]):
                cell.number_format = "yyyy-mm-dd"
            elif looks_numeric(cell.value):
                cell.number_format = '#,##0.00' if isinstance(cell.value, float) and cell.value % 1 else '#,##0'


def style_sheet(ws):
    if ws.max_row < 1 or ws.max_column < 1:
        return

    ws.sheet_view.showGridLines = True
    ws.freeze_panes = None
    ws.sheet_view.zoomScale = 95

    header = header_row(ws)
    ws.freeze_panes = ws.cell(header + 1, 1).coordinate

    for cell in ws[header]:
        if cell.value not in (None, ""):
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
            cell.font = Font(name="Arial", bold=True, color=HEADER_FONT)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = BORDER

    for row in ws.iter_rows(min_row=header + 1, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BODY_BORDER

    set_number_formats(ws, header)

    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for row in range(1, min(ws.max_row, 500) + 1):
            value = ws.cell(row, col).value
            if value is not None:
                max_len = max(max_len, len(str(value)))
        ws.column_dimensions[letter].width = min(45, max(10, max_len + 2))

    ws.row_dimensions[header].height = 24
    ref = f"A{header}:{get_column_letter(ws.max_column)}{ws.max_row}"
    if ws.max_row > header and ws.max_column > 1 and not ws.tables:
        table_name = "Table_" + "".join(ch for ch in ws.title if ch.isalnum())[:20]
        # openpyxl requires table name to be unique within workbook
        table = Table(displayName=table_name or "Table1", ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
        ws.add_table(table)

    ws.auto_filter.ref = ref
    ws.page_setup.orientation = "landscape" if ws.max_column > 8 else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_margins.left = 0.3
    ws.page_margins.right = 0.3
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.5


def main():
    if len(sys.argv) not in (2, 3):
        raise SystemExit("Usage: python scripts/style_workbook.py input.xlsx [output.xlsx]")
    src = Path(sys.argv[1])
    dst = Path(sys.argv[2]) if len(sys.argv) == 3 else src
    wb = load_workbook(src)
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1 and wb["Sheet"].max_row == 1 and wb["Sheet"].max_column == 1 and wb["Sheet"]["A1"].value is None:
        del wb["Sheet"]
    for ws in wb.worksheets:
        style_sheet(ws)
    wb.save(dst)
    print(dst)


if __name__ == "__main__":
    main()
